# -*- coding: utf-8 -*-
"""
Nạp dữ liệu chấm công (Văn phòng Hà Nội) từ các file "Cham cong HN_T*.xlsx"
vào bảng attendance_records (theo NGÀY) trên Supabase của UNI HR App.

Script chỉ nạp DỮ LIỆU THÔ theo ngày:
    check_in   = Vào 1
    check_out  = giờ ra cuối cùng (max của Ra 1, Ra 2)
    work_hours = cột "Tổng giờ" (đã trừ nghỉ trưa)
Các chỉ số đi muộn / OT / tổng hợp được TÍNH Ở APP
(lib/nhan-su/attendance.ts) để chỉ có một nguồn quy tắc duy nhất.

Chạy mỗi tháng khi có thêm dữ liệu:
    python cap_nhat_hieu_suat.py --dry-run     # chỉ in, không ghi DB
    python cap_nhat_hieu_suat.py               # nạp lên Supabase (upsert theo ngày)
"""
import openpyxl, os, sys, glob, json, urllib.request, urllib.error
from collections import defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
ENV_PATH = os.path.join(HERE, "..", "HR APP", "uni-hr-app", ".env.local")

# Mã chấm công (Mã NV MCC dạng số) chưa có trong sheet Source -> mã chuẩn
MCC_OVERRIDE = {"10": "HN00023"}   # Trịnh Thành Đạt (DAT LEGAL)

DRY = "--dry-run" in sys.argv


def load_env(path):
    env = {}
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if "=" in line and not line.startswith("#"):
                k, v = line.split("=", 1)
                env[k.strip()] = v.strip().strip('"').strip("'")
    return env


def to_min(s):
    if s is None:
        return None
    s = str(s).strip()
    if not s or ":" not in s:
        return None
    try:
        h, m = s.split(":")[:2]
        return int(h) * 60 + int(m)
    except Exception:
        return None


def min_to_time(v):
    if v is None:
        return None
    return f"{v // 60:02d}:{v % 60:02d}:00"


def norm_code(c):
    if c is None:
        return ""
    return "".join(ch for ch in str(c).strip().upper() if ch.isalnum())


def parse_date(datestr):
    """'02/01/2026' -> '2026-01-02'."""
    if datestr is None:
        return None
    s = str(datestr).strip().split(" ")[0]
    try:
        d, m, y = s.split("/")
        return f"{int(y):04d}-{int(m):02d}-{int(d):02d}"
    except Exception:
        return None


def build_source_map():
    mcc2std, std2name = dict(MCC_OVERRIDE), {}
    for f in glob.glob(os.path.join(HERE, "Cham cong HN_T*.xlsx")):
        try:
            wb = openpyxl.load_workbook(f, data_only=True)
        except Exception:
            continue
        if "Source" not in wb.sheetnames:
            continue
        ws = wb["Source"]
        for r in range(5, ws.max_row + 1):
            mcc, std, name = ws.cell(r, 3).value, ws.cell(r, 4).value, ws.cell(r, 6).value
            if std:
                std = str(std).strip()
                if mcc is not None and str(mcc).strip():
                    mcc2std.setdefault(str(mcc).strip(), std)
                if name:
                    std2name.setdefault(std, str(name).strip())
    return mcc2std, std2name


def resolve_std(raw, mcc2std):
    c = str(raw).strip()
    if c.upper().startswith("HN"):
        return c
    return mcc2std.get(c, c)


def main():
    env = load_env(ENV_PATH)
    url = env["NEXT_PUBLIC_SUPABASE_URL"]
    key = env["SUPABASE_SERVICE_ROLE_KEY"]

    mcc2std, std2name = build_source_map()

    req = urllib.request.Request(
        url + "/rest/v1/employees?select=id,full_name,staff_code",
        headers={"apikey": key, "Authorization": "Bearer " + key})
    employees = json.load(urllib.request.urlopen(req))
    code2id, name2id, id2name = {}, {}, {}
    for e in employees:
        if e.get("staff_code"):
            code2id[norm_code(e["staff_code"])] = e["id"]
        name2id[e["full_name"].strip().upper()] = e["id"]
        id2name[e["id"]] = e["full_name"]

    # (employee_id, record_date) -> row  (khử trùng lặp nếu 1 ngày xuất hiện nhiều lần)
    rows = {}
    unmatched = defaultdict(int)
    per_emp_days = defaultdict(int)

    files = sorted(glob.glob(os.path.join(HERE, "Cham cong HN_T*.xlsx")))
    print(f"Tìm thấy {len(files)} file chấm công.")
    for f in files:
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb.worksheets[0]
        for r in range(5, ws.max_row + 1):
            code = ws.cell(r, 3).value
            name = ws.cell(r, 4).value
            if not code:
                continue
            rec_date = parse_date(ws.cell(r, 1).value)
            if not rec_date:
                continue
            in1 = to_min(ws.cell(r, 5).value)
            out1 = to_min(ws.cell(r, 6).value)
            in2 = to_min(ws.cell(r, 7).value)
            out2 = to_min(ws.cell(r, 8).value)
            total = ws.cell(r, 9).value
            if in1 is None and out1 is None and in2 is None and out2 is None and not total:
                continue

            std = resolve_std(code, mcc2std)
            emp_id = code2id.get(norm_code(std))
            if not emp_id and name:
                emp_id = name2id.get(str(name).strip().upper())
            if not emp_id:
                nm = std2name.get(std)
                if nm:
                    emp_id = name2id.get(nm.strip().upper())
            if not emp_id:
                unmatched[f"{std} / {std2name.get(std, name)}"] += 1
                continue

            outs = [x for x in (out1, out2) if x is not None]
            last_out = max(outs) if outs else None
            try:
                wh = round(float(total), 2) if total is not None else None
            except Exception:
                wh = None

            rows[(emp_id, rec_date)] = {
                "employee_id": emp_id,
                "record_date": rec_date,
                "check_in": min_to_time(in1),
                "check_out": min_to_time(last_out),
                "work_hours": wh,
            }
            per_emp_days[emp_id] += 1

    payload = list(rows.values())

    print(f"\nSố ngày chấm công theo nhân viên:")
    for emp_id, n in sorted(per_emp_days.items(), key=lambda kv: id2name.get(kv[0], "")):
        nm = id2name.get(emp_id, "?")
        nm = nm.encode("ascii", "replace").decode() if os.name == "nt" else nm
        print(f"  {nm[:28]:28s} {n:4d} ngày")
    if unmatched:
        print("\n[!] Không map được nhân viên (bỏ qua):")
        for k, v in unmatched.items():
            print(f"    {k}: {v} dòng")
    print(f"\nTổng {len(payload)} bản ghi ngày (đã khử trùng lặp).")

    if DRY:
        print("\n[DRY-RUN] Không ghi vào Supabase.")
        return

    # Upsert theo lô, on_conflict = employee_id,record_date
    endpoint = url + "/rest/v1/attendance_records?on_conflict=employee_id,record_date"
    B = 500
    for i in range(0, len(payload), B):
        batch = payload[i:i + B]
        req = urllib.request.Request(
            endpoint, data=json.dumps(batch).encode("utf-8"), method="POST",
            headers={
                "apikey": key, "Authorization": "Bearer " + key,
                "Content-Type": "application/json",
                "Prefer": "resolution=merge-duplicates,return=minimal",
            })
        try:
            resp = urllib.request.urlopen(req)
            print(f"[OK] Upsert lô {i//B + 1}: {len(batch)} bản ghi (HTTP {resp.status}).")
        except urllib.error.HTTPError as e:
            print(f"[LỖI] HTTP {e.code}: {e.read().decode('utf-8', 'replace')[:300]}")
            sys.exit(1)
    print(f"[HOÀN TẤT] Đã nạp {len(payload)} bản ghi vào attendance_records.")


if __name__ == "__main__":
    main()
